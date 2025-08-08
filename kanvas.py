import sys
import os
import sqlite3
import shutil
import re
import traceback
import logging
from pathlib import Path
from datetime import datetime
from collections import defaultdict
import openpyxl
from PySide6 import __version__
from PySide6.QtWidgets import (
    QApplication, QMessageBox, QPushButton, QMainWindow, QFileDialog, QTreeView, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, 
    QComboBox, QGridLayout, QTextEdit, QSizePolicy, QWidget, QDateEdit, QProgressBar, QScrollArea, QHeaderView, QSplashScreen
)
from PySide6.QtGui import QStandardItemModel, QStandardItem, QColor, QFont, QPixmap
from PySide6.QtCore import QFile, Qt, QDate, QRect, QTimer
from viz_network import visualize_network
from viz_timeline import open_timeline_window
from helper.database_utils import create_all_tables
from helper.download_updates import download_updates
from helper.mapping_defend import open_defend_window
from helper.mapping_attack import mitre_mapping
from helper.mapping_veris import open_veris_window
from helper.bookmarks import display_bookmarks_kb
from helper.lookup_eventid import display_event_id_kb
from helper.lookup_entraid import open_entra_lookup_window
from helper.lookup_domain import open_domain_lookup_window
from helper.lookup_cve import open_cve_window
from helper.lookup_ip import open_ip_lookup_window
from helper.lookup_file import open_hash_lookup_window
from helper.lookup_ransomware import open_ransomware_kb_window
from markdown_editor import handle_markdown_editor
from filelock import FileLock, Timeout
from PySide6.QtGui import QIcon
from shiboken6 import isValid 
from helper.windowsui import Ui_KanvasMainWindow

class MainApp:
    def __init__(self):
        self.app = QApplication(sys.argv)
        logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s', filename='kanvas.log')
        self.logger = logging.getLogger(__name__)
        image_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "images")
        logo_path = os.path.join(image_dir, "logo.png")
        app_icon = QIcon(logo_path)
        self.app.setWindowIcon(app_icon)
        splash_pixmap = QPixmap(logo_path)
        if not splash_pixmap.isNull():
            splash_pixmap = splash_pixmap.scaled(300, 150, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            self.splash = QSplashScreen(splash_pixmap)
        else:
            splash_pixmap = QPixmap(300, 150)
            splash_pixmap.fill(QColor(40, 44, 52))
            self.splash = QSplashScreen(splash_pixmap)
            self.splash.showMessage("KANVAS\n \nLoading...", Qt.AlignCenter, QColor(255, 255, 255))
        self.splash.show()
        self.app.processEvents()
        self.child_windows = []
        self.db_path = "kanvas.db"
        self.file_lock = None 
        self.read_only_mode = False  
        self.splash.showMessage("Initializing database...", Qt.AlignBottom | Qt.AlignCenter, QColor(255, 255, 255))
        self.app.processEvents()
        create_all_tables(self.db_path)
        self.splash.showMessage("Loading user interface...", Qt.AlignBottom | Qt.AlignCenter, QColor(255, 255, 255))
        self.app.processEvents()
        self.window = self.load_ui()
        self.window.closeEvent = self.closeEvent
        self.current_workbook = None
        self.current_file_path = None
        self.current_sheet_name = None
        self.splash.showMessage("Connecting UI elements...", Qt.AlignBottom | Qt.AlignCenter, QColor(255, 255, 255))
        self.app.processEvents()
        self.connect_ui_elements()
        QTimer.singleShot(1000, self.finish_loading)

    def finish_loading(self):
        self.window.showMaximized()
        self.splash.finish(self.window)
    
    def get_lock_path(self, excel_path):
        return f"{excel_path}.lock"

    def acquire_file_lock(self, file_path):
        try:
            if self.file_lock:
                self.release_file_lock()  
            lock_path = self.get_lock_path(file_path)
            self.file_lock = FileLock(lock_path, timeout=1) 
            try:
                self.file_lock.acquire(timeout=1)
                self.read_only_mode = False
                return True
            except Timeout:
                self.logger.info(f"File {file_path} is locked by another process")
                result = QMessageBox.question( self.window, "File is in use", f"The file is currently being edited by another user.\nDo you want to open it in read-only mode?", QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes)
                if result == QMessageBox.Yes:
                    self.read_only_mode = True
                    return False
                else:
                    return None  
        except Exception as e:
            self.logger.error(f"Error acquiring lock: {e}")
            self.file_lock = None
            return None  
    
    def release_file_lock(self):
        if self.file_lock:
            try:
                if hasattr(self.file_lock, 'is_locked') and self.file_lock.is_locked:
                    self.file_lock.release()
                    self.logger.info("File lock released")
            except Exception as e:
                self.logger.error(f"Error releasing file lock: {e}")
            finally:
                self.file_lock = None

    def closeEvent(self, event):
        self.release_file_lock()
        for window in self.child_windows:
            if window and hasattr(window, 'setEnabled'):
                window.setEnabled(False)
        windows_to_close = self.child_windows.copy()  
        for i, window in enumerate(windows_to_close):
            try:
                if window and hasattr(window, 'close'):
                    #print(f"Closing child window {i+1}/{len(windows_to_close)}: {window.windowTitle() if hasattr(window, 'windowTitle') else 'Unknown'}")
                    window.close()
                    QApplication.processEvents()  
            except Exception as e:
                self.logger.error(f"Error during first close attempt: {e}")
        remaining = [w for w in self.child_windows if hasattr(w, 'isVisible') and w.isVisible()]
        if remaining:
            self.logger.info(f"Found {len(remaining)} windows still open, forcing deletion...")
            for window in remaining:
                try:
                    self.logger.info(f"Force closing: {window.windowTitle() if hasattr(window, 'windowTitle') else 'Unknown'}")
                    window.setAttribute(Qt.WA_DeleteOnClose, True) 
                    window.setParent(None)  
                    window.close()
                    window.deleteLater()
                    QApplication.processEvents() 
                except Exception as e:
                    self.logger.error(f"Error during forced close: {e}")
        self.child_windows.clear()
        event.accept()

    def load_ui(self):
        window = QMainWindow()
        ui = Ui_KanvasMainWindow()
        ui.setupUi(window)
        window.ui = ui
        return window
    
    def connect_ui_elements(self):
        self.tree_view = self.window.ui.treeViewMain
        self.tree_view.setSortingEnabled(True)
        if self.tree_view:
            self.tree_view.doubleClicked.connect(self.edit_row)
        else:
            self.logger.warning("treeViewMain not found!")
        self.connect_button("left_button_2", self.handle_veris_window)
        self.connect_button("left_button_3", self.handle_defend_mapping)  
        self.connect_button("left_button_4", self.handle_mitre_mapping)
        self.connect_button("left_button_5", self.handle_visualize_network)
        self.connect_button("left_button_6", self.handle_timeline_window)
        self.connect_button("left_button_7", self.open_new_case_window)
        self.connect_button("left_button_8", self.load_data_into_treeview)
        self.connect_button("left_button_9", self.handle_ip_lookup)  
        self.connect_button("left_button_10", self.handle_ransomware_kb)
        self.connect_button("left_button_11", self.handle_domain_lookup) 
        self.connect_button("left_button_12", self.handle_hash_lookup) 
        self.connect_button("left_button_15", self.handle_markdown_editor)
        self.connect_button("left_button_16", self.display_bookmarks_kb)
        self.connect_button("left_button_18", self.open_api_settings)
        self.connect_button("left_button_19", self.handle_download_updates)
        self.connect_button("left_button_17", self.display_event_id_kb)
        self.connect_button("left_button_13", self.entra_appid)
        self.connect_button("left_button_14", self.handle_cve_lookup)
        self.connect_button("down_button_1", self.add_new_row)  
        self.connect_button("down_button_2", self.delete_row)   
        self.connect_button("down_button_3", self.load_sheet)   
        self.connect_button("down_button_4", self.list_systems)  
        self.connect_button("down_button_5", self.list_users)    
        self.connect_button("down_button_6", self.sanitize)     
        
    def connect_button(self, button_name, handler):
        button = getattr(self.window.ui, button_name, None)
        if button:
            button.clicked.connect(handler)
        else:
            self.logger.warning(f"{button_name} not found!")
            
    def check_excel_loaded(self):
        if not self.current_workbook or not self.current_file_path:
            QMessageBox.warning(self.window, "Warning", "No Excel file loaded. Please load a file first.")
            return False
        return True

    def handle_veris_window(self):
        if self.check_excel_loaded():
            window = open_veris_window(self.window)
            self.track_child_window(window)

    def handle_mitre_mapping(self):
        if self.check_excel_loaded():
            window = mitre_mapping(self.window)
            self.track_child_window(window)

    def handle_visualize_network(self):
        if self.check_excel_loaded():
            window = visualize_network(self.window)
            self.track_child_window(window)

    def handle_timeline_window(self):
        if self.check_excel_loaded():
            window = open_timeline_window(self.window)
            self.track_child_window(window)
            
    def handle_download_updates(self):
        window = download_updates(self.window)
        self.track_child_window(window)

    def open_custom_window(self):
        custom_window = QWidget(self.window) 
        custom_window.setWindowTitle("New Case Details")
        custom_window.setMinimumSize(600, 800)
        layout = QVBoxLayout()
        text_box = QTextEdit()
        layout.addWidget(text_box)
        submit_button = QPushButton("Submit")
        layout.addWidget(submit_button)

        def submit_data():
            data = text_box.toPlainText()
            custom_window.close()

        submit_button.clicked.connect(submit_data)
        custom_window.setLayout(layout)
        self.track_child_window(custom_window)
        custom_window.show()
    
    def open_api_settings(self):
        custom_window = QWidget(self.window)  
        custom_window.setWindowTitle("API Settings")
        custom_window.setMinimumSize(800, 400)
        custom_window.setWindowFlags(Qt.Window | Qt.WindowTitleHint | Qt.WindowCloseButtonHint)
        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(20, 20, 20, 20)
        title_label = QLabel("API Keys Configuration")
        title_font = QFont()
        title_font.setPointSize(12)
        title_font.setBold(True)
        title_label.setFont(title_font)
        main_layout.addWidget(title_label)
        main_layout.addSpacing(10)
        
        current_settings = {}
        
        def mask_api_key(api_key):
            if not api_key or api_key.strip() == "":
                return ""
            if len(api_key) <= 8:
                return "****"
            return api_key[:4] + "*" * (len(api_key) - 4)
        
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT VT_API_KEY, SHODEN_API_KEY, OTX_API_KEY, MISP_API_KEY, OpenCTI_API_KEY, IPQS_API_KEY FROM user_settings WHERE id = 1")
            row = cursor.fetchone()
            if row:
                field_names = ['VT_API_KEY', 'SHODEN_API_KEY', 'OTX_API_KEY', 'MISP_API_KEY', 'OpenCTI_API_KEY', 'IPQS_API_KEY']
                for i, name in enumerate(field_names):
                    current_settings[name] = {
                        'original': row[i] if row[i] is not None else "",
                        'masked': mask_api_key(row[i] if row[i] is not None else "")
                    }
            else:
                self.logger.info("No settings found in database")
                for field in [
                    'VT_API_KEY', 'SHODEN_API_KEY', 'OTX_API_KEY', 'MISP_API_KEY',
                    'OpenCTI_API_KEY', 'IPQS_API_KEY'
                ]:
                    current_settings[field] = {'original': "", 'masked': ""}
        except sqlite3.Error as e:
            self.logger.error(f"Database error: {e}")
            QMessageBox.critical(custom_window, "Error", f"Failed to fetch settings: {e}")
            for field in [
                'VT_API_KEY', 'SHODEN_API_KEY', 'OTX_API_KEY', 'MISP_API_KEY',
                'OpenCTI_API_KEY', 'IPQS_API_KEY'
            ]:
                current_settings[field] = {'original': "", 'masked': ""}
        finally:
            if 'conn' in locals():
                conn.close()
        
        scroll_widget = QWidget()
        scroll_layout = QGridLayout(scroll_widget)
        scroll_layout.setColumnStretch(1, 1)  
        api_fields = [
            ("VirusTotal :", "VT_API_KEY"),
            ("Shodan.io :", "SHODEN_API_KEY"),
            ("AlienVault OTX :", "OTX_API_KEY"),
           # ("MISP :", "MISP_API_KEY"),
           # ("OpenCTI :", "OpenCTI_API_KEY"),
            ("IP Quality Score :", "IPQS_API_KEY")
        ]
        input_fields = {}
        show_buttons = {}
        for row, (label_text, field_name) in enumerate(api_fields):
            label = QLabel(label_text)
            label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            scroll_layout.addWidget(label, row, 0)
            input_field = QLineEdit()
            input_field.setMinimumWidth(350)
            input_field.setEchoMode(QLineEdit.Password) 
            original_value = current_settings.get(field_name, {}).get('original', '')
            if original_value:
                input_field.setText(original_value)
                input_field.setPlaceholderText(current_settings[field_name]['masked'])
            scroll_layout.addWidget(input_field, row, 1)
            show_button = QPushButton("Show")
            show_button.setFixedWidth(60)
            show_button.setStyleSheet("padding: 4px 8px;")
            
            def create_toggle_function(field, button):
                def toggle_visibility():
                    if field.echoMode() == QLineEdit.Password:
                        field.setEchoMode(QLineEdit.Normal)
                        button.setText("Hide")
                    else:
                        field.setEchoMode(QLineEdit.Password)
                        button.setText("Show")
                return toggle_visibility
            
            show_button.clicked.connect(create_toggle_function(input_field, show_button))
            scroll_layout.addWidget(show_button, row, 2)
            input_fields[field_name] = input_field
            show_buttons[field_name] = show_button
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setWidget(scroll_widget)
        main_layout.addWidget(scroll_area)
        #info_label = QLabel("Note: API keys are masked for security. Click 'Show' to reveal values when editing.")
        #info_label.setStyleSheet("color: #666; font-style: italic; margin: 10px 0;")
        #main_layout.addWidget(info_label)
        button_layout = QHBoxLayout()
        button_layout.addStretch(1)
        update_button = QPushButton("Update")
        update_button.setStyleSheet("background-color: #4CAF50; color: white; padding: 6px 20px;")
        button_layout.addWidget(update_button)
        cancel_button = QPushButton("Cancel")
        cancel_button.setStyleSheet("padding: 6px 20px;")
        button_layout.addWidget(cancel_button)
        main_layout.addSpacing(15)
        main_layout.addLayout(button_layout)
        custom_window.setLayout(main_layout)

        def save_settings():
            try:
                conn = sqlite3.connect(self.db_path)
                cursor = conn.cursor()
                update_fields = []
                values = []
                for field_name, input_field in input_fields.items():
                    update_fields.append(f"{field_name} = ?")
                    api_key_value = input_field.text().strip()
                    values.append(api_key_value)
                query = f"UPDATE user_settings SET {', '.join(update_fields)} WHERE id = 1"
                cursor.execute(query, values)
                if cursor.rowcount == 0:
                    field_names = list(input_fields.keys())
                    placeholders = ", ".join(["?"] * len(field_names))
                    field_str = ", ".join(field_names)
                    query = f"INSERT INTO user_settings (id, {field_str}) VALUES (1, {placeholders})"
                    cursor.execute(query, values)
                conn.commit()
                QMessageBox.information(custom_window, "Success", "Settings saved successfully!")
                custom_window.close()
            except sqlite3.Error as e:
                self.logger.error(f"Error saving settings: {e}")
                QMessageBox.critical(custom_window, "Error", f"Failed to save settings: {e}")
            finally:
                if 'conn' in locals():
                    conn.close()
    
        update_button.clicked.connect(save_settings)
        cancel_button.clicked.connect(custom_window.close)
        self.child_windows.append(custom_window)
        custom_window.show()

    def open_new_case_window(self):
        file_path, _ = QFileDialog.getSaveFileName(self.window, "Save New Case File", "New_Case.xlsx", "Excel files (*.xlsx)")
        if not file_path:
            QMessageBox.warning(self.window, "Warning", "No file selected. Case creation canceled.")
            return
        lock_acquired = False
        try:
            lock_status = self.acquire_file_lock(file_path)
            if lock_status is None:  
                return
            lock_acquired = True
            template_file = "sod.xlsx"
            if not os.path.exists(template_file):
                QMessageBox.critical(self.window, "Error", f"Template file '{template_file}' not found.")
                return
            shutil.copy(template_file, file_path)
            QMessageBox.information(self.window, "Success", f"New case created successfully and saved to {file_path}!")
            workbook = openpyxl.load_workbook(file_path)
            self.current_file_path = file_path
            self.current_workbook = workbook
            self.read_only_mode = False  
            file_status_label = self.window.ui.labelFileStatus
            if file_status_label:
                file_name = os.path.basename(file_path)
                file_status_label.setText(f"Loaded: {file_name} ({file_path})")
            else:
                self.logger.warning("labelFileStatus not found!")
            if self.tree_view:
                self.load_data_into_treeview(file_path=file_path, workbook=workbook)
            else:
                self.logger.warning("treeViewMain not found!")
        except Exception as e:
            self.logger.error(f"Error creating new case: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to create new case: {e}")
        finally:
            if lock_acquired and (not hasattr(self, 'current_workbook') or not self.current_workbook):
                self.release_file_lock()

    def load_data_into_treeview(self, file_path=None, workbook=None):
        lock_acquired = False
        progress = None
        try:
            if file_path is None or workbook is None:
                file_path, _ = QFileDialog.getOpenFileName(
                    self.window, "Open Excel File", "", "Excel Files (*.xlsx *.xls)"
                )
                if not file_path:
                    self.logger.info("No file selected")
                    return
                progress = QProgressBar()
                progress.setWindowTitle("Loading Excel File")
                progress.setGeometry(QRect(300, 300, 400, 30))
                progress.setWindowFlags(Qt.Window | Qt.WindowStaysOnTopHint)
                progress.setRange(0, 0)  
                progress.show()
                progress.setFormat("Loading file...")
                QApplication.processEvents()
                try:
                    lock_status = self.acquire_file_lock(file_path)
                    if lock_status is None:  
                        if progress:
                            progress.close()
                        return
                    lock_acquired = True
                    progress.setFormat("Opening workbook...")
                    QApplication.processEvents()
                    workbook = openpyxl.load_workbook(file_path, read_only=self.read_only_mode)
                except Exception as e:
                    if progress:
                        progress.close()
                    if lock_acquired:
                        self.release_file_lock()
                        lock_acquired = False
                    self.logger.error(f"Error loading Excel file: {e}")
                    QMessageBox.critical(self.window, "Error", f"Failed to load Excel file: {e}")
                    return

            self.current_workbook = workbook
            self.current_file_path = file_path
            self.window.current_workbook = workbook  
            self.window.current_file_path = file_path  
            if progress:
                progress.setFormat("Updating file status...")
                QApplication.processEvents()
            file_status_label = self.window.ui.labelFileStatus
            if file_status_label:
                file_name = os.path.basename(file_path)
                read_only_text = " [READ-ONLY]" if self.read_only_mode else ""
                file_status_label.setText(f"Loaded: {file_name} ({file_path}){read_only_text}")
            else:
                self.logger.warning("labelFileStatus not found!")
            sheet_dropdown = self.window.ui.comboBoxSheet
            if not sheet_dropdown:
                if progress:
                    progress.close()
                self.logger.warning("Sheet dropdown (comboBoxSheet) not found!")
                return
            if progress:
                progress.setFormat("Loading sheet names...")
                QApplication.processEvents()
            sheet_dropdown.clear()
            try:
                sheet_dropdown.currentIndexChanged.disconnect()
            except:
                pass  # No connections to disconnect
            sheet_dropdown.addItems(workbook.sheetnames)
        
            def load_selected_sheet():
                selected_sheet_name = sheet_dropdown.currentText()
                if not selected_sheet_name:
                    self.logger.info("No sheet selected")
                    return
                sheet_progress = QProgressBar()
                sheet_progress.setWindowTitle(f"Loading Sheet: {selected_sheet_name}")
                sheet_progress.setGeometry(QRect(300, 300, 400, 30))
                sheet_progress.setWindowFlags(Qt.Window | Qt.WindowStaysOnTopHint)
                sheet_progress.show()
                try:
                    current_workbook = self.current_workbook
                    if not current_workbook:
                        sheet_progress.close()
                        return
                    self.current_sheet_name = selected_sheet_name
                    self.window.current_sheet_name = selected_sheet_name  
                    sheet = current_workbook[selected_sheet_name]
                    max_row = sheet.max_row
                    sheet_progress.setRange(0, max_row + 2)  
                    sheet_progress.setValue(0)
                    sheet_progress.setFormat("Reading headers...")
                    QApplication.processEvents()
                    model = QStandardItemModel()
                    headers = [cell.value for cell in sheet[1]]  
                    model.setHorizontalHeaderLabels(headers)
                    sheet_progress.setValue(1)
                    sheet_progress.setFormat("Loading data...")
                    QApplication.processEvents()
                    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):  
                        items = [QStandardItem(str(cell.value) if cell.value is not None else "") for cell in row]
                        model.appendRow(items)
                        if row_index % 100 == 0 or row_index == max_row:
                            sheet_progress.setValue(row_index)
                            sheet_progress.setFormat(f"Loading row {row_index}/{max_row}...")
                            QApplication.processEvents()
                    sheet_progress.setFormat("Finalizing...")
                    QApplication.processEvents()
                    self.tree_view.setModel(model)
                    self.tree_view.setEditTriggers( 
                        QTreeView.NoEditTriggers if self.read_only_mode else 
                        (QTreeView.DoubleClicked | QTreeView.EditKeyPressed | QTreeView.AnyKeyPressed)
                    )
                    sheet_progress.setValue(max_row + 2)
                    sheet_progress.close()
                except Exception as e:
                    sheet_progress.close()
                    self.logger.error(f"Error loading sheet '{selected_sheet_name}': {e}")
                    self.tree_view.setModel(QStandardItemModel())
            sheet_dropdown.currentIndexChanged.connect(load_selected_sheet)
            if progress:
                progress.setFormat("Initializing sheet view...")
                QApplication.processEvents()
            sheet_dropdown.setCurrentIndex(0)
            load_selected_sheet()
            if progress:
                progress.close()
        except Exception as e:
            if progress:
                progress.close()
            self.logger.error(f"Error in load_data_into_treeview: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to load data into tree view: {e}")
            if lock_acquired and not workbook:  
                self.release_file_lock()
                
    def get_evidence_types_from_db(self):
        evidence_types = []
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT DISTINCT evidencetype FROM EvidenceType ORDER BY evidencetype")
            evidence_types = [row[0] for row in cursor.fetchall()]
            conn.close()
            if not evidence_types:
                evidence_types = []
        except sqlite3.Error as e:
            self.logger.error(f"Error fetching evidence types from database: {e}")
            evidence_types = []
        return evidence_types

    def edit_row(self, index):
        if self.read_only_mode:
            QMessageBox.information(self.window, "Read-Only Mode", "This file is open in read-only mode. You cannot edit its contents.")
            return
        if not index.isValid():
            QMessageBox.warning(self.window, "Warning", "No row selected. Please select a row to edit.")
            return
        row_index = index.row()
        model = self.tree_view.model()
        if not model:
            QMessageBox.critical(self.window, "Error", "No model found for the tree view.")
            return
        mitre_tactic_options = []
        mitre_techniques_by_tactic = {}  
        current_tactic = None
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT DISTINCT PID FROM mitre_techniques")
            mitre_tactic_options = [row[0] for row in cursor.fetchall()]
            for tactic in mitre_tactic_options:
                cursor.execute("SELECT ID FROM mitre_techniques WHERE PID = ?", (tactic,))
                mitre_techniques_by_tactic[tactic] = [row[0] for row in cursor.fetchall()]
            conn.close()
        except sqlite3.Error as e:
            QMessageBox.critical(self.window, "Error", f"Failed to fetch MITRE values: {e}")
        evidence_types = self.get_evidence_types_from_db()
        editor_widget = QWidget(self.window)
        editor_widget.setWindowTitle("Edit Row - Widget Editor")
        editor_widget.setMinimumSize(500, 400)
        editor_widget.setWindowFlags(Qt.Window)  
        grid_layout = QGridLayout()
        grid_layout.setContentsMargins(15, 15, 15, 15)
        grid_layout.setHorizontalSpacing(20)
        grid_layout.setVerticalSpacing(10)
        input_fields = []
        mitre_tactic_combo = None
        mitre_technique_combo = None
        existing_tactic_value = None
        existing_technique_value = None
        for col in range(model.columnCount()):
            header_text = model.headerData(col, Qt.Horizontal)
            cell_data = model.index(row_index, col).data()
            header_label = QLabel(f"{header_text}:")
            header_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            if header_text and header_text.strip().lower() == "visualize":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No"])
                if cell_data in ["Yes", "No"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "evidencetype":
                combo_box = QComboBox()
                combo_box.addItems(evidence_types)
                if cell_data in evidence_types:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "indicatortype":
                combo_box = QComboBox()
                combo_box.addItems(["IPAddress","UserName","FileName","FilePath","UserAgent","DomainName","JA3-JA3S","URL","Mutex","Other-Strings","EmailAddress","RegistryPath","GPO"])
                if cell_data in ["IPAddress","UserName","FileName","FilePath","UserAgent","DomainName","JA3-JA3S","URL","Mutex","Other-Strings","EmailAddress","RegistryPath","GPO"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "location":
                combo_box = QComboBox()
                combo_box.addItems(["On-Prem", "Unknown", "Cloud-Generic", "Cloud-Azure", "Cloud-AWS", "Clous-GCP"])
                if cell_data in ["On-Prem", "Unknown", "Cloud-Generic", "Cloud-Azure", "Cloud-AWS", "Clous-GCP"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "currentstatus":
                combo_box = QComboBox()
                combo_box.addItems(["Completed" , "In Progress", "On Hold", "Not Started"  ])
                if cell_data in ["Completed" , "In Progress", "On Hold", "Not Started"  ]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "priority":
                combo_box = QComboBox()
                combo_box.addItems(["High" ,"Medium" ,"Low" ])
                if cell_data in ["High" ,"Medium" ,"Low" ]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box                
            elif header_text and header_text.strip().lower() == "evidencecollected":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No" ])
                if cell_data in ["Yes", "No" ]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box                
            elif header_text and header_text.strip().lower() == "targettype":
                combo_box = QComboBox()
                combo_box.addItems(["Machine", "Identity", "Others"])
                if cell_data in ["Machine", "Identity", "Others"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box               
            elif header_text and header_text.strip().lower() == "systemtype":
                combo_box = QComboBox()
                combo_box.addItems(["Attacker-Machine","Server-Generic","Server-Application","Server-Web","Server-DC","Server-Terminal SRV","Server-Database","Gateway-Generic", "Gateway-Firewall","Gateway-VPN","Gateway-Router","Gateway-Switch","Gateway-Email","Gateway-Web Proxy","Gateway-DNS","Desktop","Mobile","OT Device","UnKnown"])
                if cell_data in ["Attacker-Machine","Server-Generic","Server-Application","Server-Web","Server-DC","Server-Terminal SRV","Server-Database","Gateway-Generic", "Gateway-Firewall","Gateway-VPN","Gateway-Router","Gateway-Switch","Gateway-Email","Gateway-Web Proxy","Gateway-DNS","Desktop","Mobile","OT Device","UnKnown"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "accounttype":
                combo_box = QComboBox()
                combo_box.addItems(["Normal User Account - Local","Normal User Account - On-Prem AD","Normal User Account - Azure","Service Account", "Domain Admin", "Global Admin - Azure", "Service Principle - Azure", "Computer Account", "Local Administrator"])
                if cell_data in ["Normal User Account - Local","Normal User Account - On-Prem AD","Normal User Account - Azure","Service Account", "Domain Admin", "Global Admin - Azure", "Service Principle - Azure", "Computer Account", "Local Administrator"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text.strip().lower() == "entrypoint":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No" ])
                if cell_data in ["Yes", "No" ]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box               
            elif header_text and (header_text.strip().lower() == "notes" or header_text.strip().lower() == "activity"):
                text_edit = QTextEdit()
                text_edit.setPlainText(cell_data if cell_data else "")
                input_field = text_edit
            elif header_text and header_text.strip().lower() in ["date added", "date updated", "date completed", "date requested", "date received"]:
                date_edit = QDateEdit()
                date_edit.setCalendarPopup(True)  
                if cell_data:
                    try:
                        date_edit.setDate(QDate.fromString(cell_data, "yyyy-MM-dd"))  
                    except Exception as e:
                        self.logger.error(f"Error parsing date: {e}")
                input_field = date_edit
            elif header_text and header_text == "TLP":
                combo_box = QComboBox()
                combo_box.addItems(["TLP-Red", "TLP-Amber_Strict", "TLP-Amber", "TLP-Green", "TLP-Clear"])
                if cell_data in ["TLP-Red", "TLP-Amber_Strict", "TLP-Amber", "TLP-Green", "TLP-Clear"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            elif header_text and header_text == "MITRE Tactic":
                existing_tactic_value = cell_data if cell_data else ""
                combo_box = QComboBox()
                combo_box.addItem("")  
                combo_box.addItems(mitre_tactic_options)
                if existing_tactic_value and existing_tactic_value in mitre_tactic_options:
                    combo_box.setCurrentText(existing_tactic_value)
                mitre_tactic_combo = combo_box
                input_field = combo_box
            elif header_text and header_text == "MITRE Techniques":
                existing_technique_value = cell_data if cell_data else ""
                combo_box = QComboBox()
                combo_box.addItem("")
                mitre_technique_combo = combo_box
                input_field = combo_box
            elif header_text and header_text == "<->":
                combo_box = QComboBox()
                combo_box.addItems([" ","->", "<-", "<->"])
                if cell_data in [" ","->", "<-", "<->"]:
                    combo_box.setCurrentText(cell_data)
                input_field = combo_box
            else:
                line_edit = QLineEdit()
                line_edit.setText(cell_data if cell_data else "")
                line_edit.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
                input_field = line_edit
            grid_layout.addWidget(header_label, col, 0)
            grid_layout.addWidget(input_field, col, 1)
            input_fields.append(input_field)
        if mitre_tactic_combo and mitre_technique_combo:
            current_tactic = mitre_tactic_combo.currentText()
            if current_tactic and current_tactic in mitre_techniques_by_tactic:
                available_techniques = mitre_techniques_by_tactic[current_tactic]
                mitre_technique_combo.addItems(available_techniques)
                if existing_technique_value:
                    if existing_technique_value in available_techniques:
                        mitre_technique_combo.setCurrentText(existing_technique_value)
                    else:
                        mitre_technique_combo.addItem(existing_technique_value)
                        mitre_technique_combo.setCurrentText(existing_technique_value)

            def update_techniques(index):
                selected_tactic = mitre_tactic_combo.currentText()
                current_technique = mitre_technique_combo.currentText()
                mitre_technique_combo.clear()
                mitre_technique_combo.addItem("")
                if selected_tactic and selected_tactic in mitre_techniques_by_tactic:
                    techniques = mitre_techniques_by_tactic[selected_tactic]
                    mitre_technique_combo.addItems(techniques)
                    if current_technique and current_technique in techniques:
                        mitre_technique_combo.setCurrentText(current_technique)
            mitre_tactic_combo.currentIndexChanged.connect(update_techniques)
        button_layout = QHBoxLayout()
        button_layout.setSpacing(20)
        save_button = QPushButton("Save")
        cancel_button = QPushButton("Cancel")
        button_layout.addStretch()
        button_layout.addWidget(save_button)
        button_layout.addWidget(cancel_button)
        button_layout.addStretch()
        main_layout = QVBoxLayout(editor_widget)
        main_layout.addLayout(grid_layout)
        main_layout.addLayout(button_layout)
        editor_widget.setLayout(main_layout)

        def save_changes():
            try:
                if not self.current_workbook or not self.current_sheet_name:
                    QMessageBox.critical(self.window, "Error", "No workbook or sheet loaded to save changes.")
                    return
                workbook = self.current_workbook
                sheet = workbook[self.current_sheet_name]
                for col, field in enumerate(input_fields):
                    if isinstance(field, QComboBox):
                        new_text = field.currentText()
                    elif isinstance(field, QTextEdit):
                        new_text = field.toPlainText()
                    elif isinstance(field, QDateEdit):
                        new_text = field.date().toString("yyyy-MM-dd")  
                    else:
                        new_text = field.text()
                    sheet.cell(row=row_index + 2, column=col + 1, value=new_text)
                    model.setData(model.index(row_index, col), new_text)
                workbook.save(self.current_file_path)
                editor_widget.close()
            except Exception as e:
                self.logger.error(f"Error in edit_row save_changes: {e}")
                QMessageBox.critical(self.window, "Error", f"Failed to save changes: {e}")
        save_button.clicked.connect(save_changes)
        cancel_button.clicked.connect(editor_widget.close)
        editor_widget.show()
        self.track_child_window(editor_widget)

    def display_bookmarks_kb(self):
        window = display_bookmarks_kb(self, self.db_path)
        self.track_child_window(window)

    def display_event_id_kb(self):
        try:
            window = display_event_id_kb(self, self.db_path)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error in display_event_id_kb: {e}")
            traceback.print_exc()
            QMessageBox.critical(self.window, "Error", f"Failed to open Event ID Knowledge Base: {str(e)}")

    def entra_appid(self):
        window = open_entra_lookup_window(self.window, self.db_path)
        self.track_child_window(window)

    def handle_cve_lookup(self):
        try:
            window = open_cve_window(self.window, self.db_path)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening CVE lookup window: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open CVE lookup window: {e}")

    def handle_ransomware_kb(self):
        try:
            window = open_ransomware_kb_window(self.window)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening Ransomware Knowledge Base window: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open Ransomware Knowledge Base window: {e}")

    def handle_defend_mapping(self):
        if self.check_excel_loaded():
            try:
                if 'Timeline' in self.current_workbook.sheetnames:
                    window = open_defend_window(self.window, self.current_file_path)
                    self.track_child_window(window)
                else:
                    QMessageBox.warning(self.window, "Missing Sheet", "The required 'Timeline' sheet was not found in this workbook.")
            except Exception as e:
                self.logger.error(f"Error in D3FEND mapping: {e}")
                QMessageBox.critical(self.window, "D3FEND Mapping Error", f"An error occurred while opening the D3FEND mapping window:\n\n{str(e)}")

    def handle_ip_lookup(self):
        try:
            window = open_ip_lookup_window(self.window, self.db_path)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening IP lookup window: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open IP lookup window: {e}")
            
    def handle_domain_lookup(self):
        try:
            window = open_domain_lookup_window(self.window, self.db_path)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening Domain lookup window: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open Domain lookup window: {e}")

    def handle_hash_lookup(self):
        try:
            window = open_hash_lookup_window(self.window, self.db_path)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening Hash lookup window: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open Hash lookup window: {e}")

    def handle_markdown_editor(self):
        try:
            window = handle_markdown_editor(self.window)
            self.track_child_window(window)
        except Exception as e:
            self.logger.error(f"Error opening Markdown Editor interface: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to open Markdown Editor interface: {e}")

    def add_new_row(self):
        if self.read_only_mode:
            QMessageBox.information(self.window, "Read-Only Mode", "This file is open in read-only mode. You cannot add new rows.")
            return
        if not self.current_sheet_name or not self.current_workbook:
            QMessageBox.warning(self.window, "Warning", "Please select a sheet first.")
            return
        mitre_tactic_options = []
        mitre_techniques_by_tactic = {}  
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT DISTINCT PID FROM mitre_techniques")
            mitre_tactic_options = [row[0] for row in cursor.fetchall()]
            for tactic in mitre_tactic_options:
                cursor.execute("SELECT ID FROM mitre_techniques WHERE PID = ?", (tactic,))
                mitre_techniques_by_tactic[tactic] = [row[0] for row in cursor.fetchall()]
            conn.close()
        except sqlite3.Error as e:
            self.logger.error(f"Error fetching MITRE values: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to fetch MITRE values: {e}")
            return
        evidence_types = self.get_evidence_types_from_db()
        sheet = self.current_workbook[self.current_sheet_name]
        headers = [cell.value for cell in sheet[1]]  
        add_row_window = QWidget(self.window)
        add_row_window.setWindowTitle("Add New Row - Widget Editor")
        add_row_window.setMinimumSize(500, 400)
        add_row_window.setWindowFlags(Qt.Window)  
        main_layout = QVBoxLayout(add_row_window)
        grid_layout = QGridLayout()
        grid_layout.setContentsMargins(15, 15, 15, 15)
        grid_layout.setHorizontalSpacing(20)
        grid_layout.setVerticalSpacing(10)
        input_fields = []
        mitre_tactic_combo = None
        mitre_technique_combo = None
        current_tactic = None
        for col, header in enumerate(headers):
            header_label = QLabel(f"{header}:")
            header_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            grid_layout.addWidget(header_label, col, 0)
            if header and header.strip().lower() == "visualize":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No"])
                input_field = combo_box
            elif header and header.strip().lower() == "evidencetype":
                combo_box = QComboBox()
                combo_box.addItems(evidence_types)
                input_field = combo_box
            elif header and header.strip().lower() == "indicatortype":
                combo_box = QComboBox()
                combo_box.addItems(["IPAddress","UserName","FileName","FilePath","UserAgent","DomainName","JA3-JA3S","URL","Mutex","Other-Strings","EmailAddress","RegistryPath","GPO"])
                input_field = combo_box
            elif header and header.strip().lower() == "accounttype":
                combo_box = QComboBox()
                combo_box.addItems(["Normal User Account - Local","Normal User Account - On-Prem AD","Normal User Account - Azure","Service Account", "Domain Admin", "Global Admin - Azure", "Service Principle - Azure", "Computer Account", "Local Administrator"])
                input_field = combo_box
            elif header and header.strip().lower() == "systemtype":
                combo_box = QComboBox()
                combo_box.addItems(["Attacker-Machine","Server-Generic","Server-Application","Server-Web","Server-DC","Server-Terminal SRV","Server-Database","Gateway-Generic", "Gateway-Firewall","Gateway-VPN","Gateway-Router","Gateway-Switch","Gateway-Email","Gateway-Web Proxy","Gateway-DNS","Desktop","Mobile","OT Device","UnKnown"])
                input_field = combo_box
            elif header and header.strip().lower() == "location":
                combo_box = QComboBox()
                combo_box.addItems(["On-Prem", "Unknown", "Cloud-Generic", "Cloud-Azure", "Cloud-AWS", "Clous-GCP"])
                input_field = combo_box
            elif header and header.strip().lower() == "currentstatus":
                combo_box = QComboBox()
                combo_box.addItems(["Completed" , "In Progress", "On Hold", "Not Started"  ])
                input_field = combo_box                
            elif header and header.strip().lower() == "priority":
                combo_box = QComboBox()
                combo_box.addItems(["High" ,"Medium" ,"Low" ])
                input_field = combo_box                
            elif header and header.strip().lower() == "evidencecollected":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No" ])
                input_field = combo_box                
            elif header and header.strip().lower() == "entrypoint":
                combo_box = QComboBox()
                combo_box.addItems(["Yes", "No" ])
                input_field = combo_box                
            elif header and header.strip().lower() == "targettype":
                combo_box = QComboBox()
                combo_box.addItems(["Machine", "Identity", "Others"])
                input_field = combo_box                
            elif header and (header.strip().lower() == "notes" or header.strip().lower() == "activity"):
                text_edit = QTextEdit()
                input_field = text_edit
            elif header and "timestamp" in header.strip().lower() and "utc" in header.strip().lower():
                line_edit = QLineEdit()
                line_edit.setPlaceholderText("YYYY-MM-DD HH:MM:SS")
                line_edit.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
                input_field = line_edit
            elif header and header.strip().lower() in ["date added", "date updated", "date completed", "date requested", "date received"]:
                date_edit = QDateEdit()
                date_edit.setCalendarPopup(True)  
                date_edit.setDate(QDate.currentDate())  
                input_field = date_edit
            elif header == "MITRE Tactic":
                combo_box = QComboBox()
                combo_box.addItem("")  
                combo_box.addItems(mitre_tactic_options)
                mitre_tactic_combo = combo_box
                input_field = combo_box
            elif header == "MITRE Techniques":
                combo_box = QComboBox()
                combo_box.addItem("")  
                mitre_technique_combo = combo_box
                input_field = combo_box
            elif header == "TLP":
                combo_box = QComboBox()
                combo_box.addItems(["TLP-Red", "TLP-Amber_Strict", "TLP-Amber", "TLP-Green", "TLP-Clear"])
                input_field = combo_box
            elif header == "<->":
                combo_box = QComboBox()
                combo_box.addItems([" ","->", "<-", "<->"])
                input_field = combo_box
            else:
                line_edit = QLineEdit()
                line_edit.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
                input_field = line_edit
            grid_layout.addWidget(header_label, col, 0)
            grid_layout.addWidget(input_field, col, 1)
            input_fields.append(input_field)
        if mitre_tactic_combo and mitre_technique_combo:
            current_technique = mitre_technique_combo.currentText()
            
            def update_techniques(tactic_index):
                current_selection = mitre_technique_combo.currentText()
                selected_tactic = mitre_tactic_combo.currentText()
                mitre_technique_combo.clear()
                mitre_technique_combo.addItem("")  
                if selected_tactic:
                    techniques = mitre_techniques_by_tactic.get(selected_tactic, [])
                    mitre_technique_combo.addItems(techniques)
                    if current_selection and current_selection in techniques:
                        mitre_technique_combo.setCurrentText(current_selection)
                else:
                    if current_selection:
                        mitre_technique_combo.addItem(current_selection)
                        mitre_technique_combo.setCurrentText(current_selection)
            mitre_tactic_combo.currentIndexChanged.connect(update_techniques)
        button_layout = QHBoxLayout()
        button_layout.setSpacing(20)
        save_button = QPushButton("Save")
        cancel_button = QPushButton("Cancel")
        button_layout.addStretch()
        button_layout.addWidget(save_button)
        button_layout.addWidget(cancel_button)
        button_layout.addStretch()
        main_layout.addLayout(grid_layout)
        main_layout.addLayout(button_layout)
    
        def save_new_row():
            try:
                if not self.current_workbook or not self.current_sheet_name:
                    QMessageBox.critical(self.window, "Error", "No workbook or sheet loaded to save changes.")
                    return
                sheet = self.current_workbook[self.current_sheet_name]
                new_row_data = []
                for field in input_fields:
                    if isinstance(field, QComboBox):
                        new_text = field.currentText()
                    elif isinstance(field, QTextEdit):
                        new_text = field.toPlainText()
                    elif isinstance(field, QDateEdit):
                        new_text = field.date().toString("yyyy-MM-dd")  
                    else:
                        new_text = field.text()
                    new_row_data.append(new_text)
                sheet.append(new_row_data)
                self.current_workbook.save(self.current_file_path)
                self.load_sheet()
                add_row_window.close()
            except Exception as e:
                self.logger.error(f"Failed to save changes: {e}")
                QMessageBox.critical(add_row_window, "Error", f"Failed to save changes: {e}")
        save_button.clicked.connect(save_new_row)
        cancel_button.clicked.connect(add_row_window.close)
        add_row_window.show()
        self.track_child_window(add_row_window)

    def delete_row(self):
        if self.read_only_mode:
            QMessageBox.information(self.window, "Read-Only Mode", "This file is open in read-only mode. You cannot delete rows.")
            return
        if not self.current_sheet_name or not self.current_workbook:
            QMessageBox.warning(self.window, "Warning", "Please select a sheet.")
            return
        if not self.tree_view or not self.tree_view.model():
            QMessageBox.warning(self.window, "Warning", "Tree view not available.")
            return
        selected_indices = self.tree_view.selectionModel().selectedRows()
        if not selected_indices:
            QMessageBox.warning(self.window, "Warning", "Please select rows to delete.")
            return
        row_indices = sorted([index.row() + 2 for index in selected_indices], reverse=True)
        confirm = QMessageBox.question(self.window, "Confirm Deletion", f"Are you sure you want to delete {len(row_indices)} row(s)?", QMessageBox.Yes | QMessageBox.No)
        if confirm == QMessageBox.No:
            return
        sheet = self.current_workbook[self.current_sheet_name]
        try:
            for row_idx in row_indices:
                sheet.delete_rows(row_idx, 1)  
            self.current_workbook.save(self.current_file_path)
            QMessageBox.information(self.window, "Success", "Rows deleted successfully!")
            self.load_sheet()
        except PermissionError:
            self.logger.error("Permission error: Cannot write to file")
            QMessageBox.critical(self.window, "Error", "Cannot save file. Ensure it is not open in another program.")
        except Exception as e:
            self.logger.error(f"Error deleting rows: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to delete rows: {e}")

    def load_sheet(self):
        if not self.current_workbook or not self.current_sheet_name or not self.tree_view:
            self.logger.error("Cannot load sheet: Missing workbook, sheet name, or tree view")
            return
        try:
            sheet = self.current_workbook[self.current_sheet_name]
            model = QStandardItemModel()
            headers = []
            for col in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col).value
                if header:
                    headers.append(str(header))
                else:
                    headers.append(f"Column {col}")
            model.setHorizontalHeaderLabels(headers)
            for row_idx in range(2, sheet.max_row + 1):
                row_data = []
                for col in range(1, sheet.max_column + 1):
                    cell_value = sheet.cell(row=row_idx, column=col).value
                    row_data.append(str(cell_value) if cell_value is not None else "")
                items = [QStandardItem(value) for value in row_data]
                model.appendRow(items)
            self.tree_view.setModel(model)
            self.tree_view.setEditTriggers(QTreeView.NoEditTriggers)
            for col_idx, header in enumerate(headers):
                self.tree_view.setColumnWidth(col_idx, 100)
        except Exception as e:
            self.logger.error(f"Error loading sheet: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to load sheet: {e}")

    def run(self):
        self.app.aboutToQuit.connect(self.application_cleanup)
        return self.app.exec()

    def list_systems(self):
        if not self.current_workbook:
            QMessageBox.warning(self.window, "Warning", "No workbook loaded. Please load a workbook first.")
            return
        if 'Timeline' not in self.current_workbook.sheetnames:
            QMessageBox.warning(self.window, "Missing Sheet", "The 'Timeline' sheet was not found in the current workbook.")
            return
        try:
            sheet = self.current_workbook['Timeline']
            event_system_col = None
            remote_system_col = None
            for col in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col).value
                if header == "Event System":
                    event_system_col = col
                elif header == "Remote System":
                    remote_system_col = col;
            if event_system_col is None and remote_system_col is None:
                QMessageBox.warning(self.window, "Missing Columns", "Neither 'Event System' nor 'Remote System' columns found in the Timeline sheet.")
                return
            systems = set()
            for row in range(2, sheet.max_row + 1): 
                if event_system_col:
                    event_system = sheet.cell(row=row, column=event_system_col).value
                    if event_system:
                        systems.add(str(event_system).strip())
                if remote_system_col:
                    remote_system = sheet.cell(row=row, column=remote_system_col).value
                    if remote_system:
                        systems.add(str(remote_system).strip())
            sorted_systems = sorted(list(systems))
            systems_window = QWidget(self.window)  
            systems_window.setWindowTitle("Unique Systems")
            systems_window.resize(400, 500)
            systems_window.setWindowFlags(Qt.Window | Qt.WindowTitleHint | Qt.WindowCloseButtonHint)
            layout = QVBoxLayout(systems_window)
            label = QLabel(f"Found {len(sorted_systems)} unique systems:")
            label.setFont(QFont("Arial", 10, QFont.Bold))
            layout.addWidget(label)
            list_widget = QTreeView()
            model = QStandardItemModel()
            model.setHorizontalHeaderLabels(["System Name"])
            for system in sorted_systems:
                item = QStandardItem(system)
                model.appendRow(item)
            list_widget.setModel(model)
            list_widget.setAlternatingRowColors(True)
            list_widget.setEditTriggers(QTreeView.NoEditTriggers)
            header = list_widget.header()
            header.setSectionResizeMode(0, QHeaderView.Stretch)
            layout.addWidget(list_widget, 1) 
            button_layout = QHBoxLayout()
            close_button = QPushButton("Close")
            close_button.clicked.connect(systems_window.close)
            button_layout.addStretch()
            button_layout.addWidget(close_button)
            button_layout.addStretch()
            layout.addLayout(button_layout)
            self.child_windows.append(systems_window)
            systems_window.show()
        except Exception as e:
            self.logger.error(f"Error listing systems: {e}")
            QMessageBox.critical(self.window, "Error", f"Failed to list systems: {e}")
        
    def close_systems_window(self):
        if hasattr(self, 'systems_window'):
            self.systems_window.close()
            self.systems_window = None

    def list_users(self):
        if not self.current_workbook:
            QMessageBox.warning(self.window, "Warning", "No workbook loaded. Please load a workbook first.")
            return
        if 'Timeline' not in self.current_workbook.sheetnames:
            QMessageBox.warning(self.window, "Missing Sheet", "The 'Timeline' sheet was not found in the current workbook.")
            return
        try:
            sheet = self.current_workbook['Timeline']
            suspect_account_col = None
            for col in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col).value
                if header == "Suspect Account":
                    suspect_account_col = col
                    break
            if suspect_account_col is None:
                QMessageBox.warning(self.window, "Missing Column", "'Suspect Account' column not found in the Timeline sheet.")
                return
            users = set()
            for row in range(2, sheet.max_row + 1):  
                user = sheet.cell(row=row, column=suspect_account_col).value
                if user:
                    users.add(str(user).strip())
            sorted_users = sorted(list(users))
            users_window = QWidget(self.window)  
            users_window.setWindowTitle("Unique Users")
            users_window.resize(400, 500)
            users_window.setWindowFlags(Qt.Window | Qt.WindowTitleHint | Qt.WindowCloseButtonHint)
            layout = QVBoxLayout(users_window)
            label = QLabel(f"Found {len(sorted_users)} unique users:")
            label.setFont(QFont("Arial", 10, QFont.Bold))
            layout.addWidget(label)
            list_widget = QTreeView()
            model = QStandardItemModel()
            model.setHorizontalHeaderLabels(["Username"])
            for user in sorted_users:
                item = QStandardItem(user)
                model.appendRow(item)
            list_widget.setModel(model)
            list_widget.setAlternatingRowColors(True)
            list_widget.setEditTriggers(QTreeView.NoEditTriggers)
            header = list_widget.header()
            header.setSectionResizeMode(0, QHeaderView.Stretch)
            layout.addWidget(list_widget, 1)  
            button_layout = QHBoxLayout()
            close_button = QPushButton("Close")
            close_button.clicked.connect(users_window.close)
            button_layout.addStretch()
            button_layout.addWidget(close_button)
            button_layout.addStretch()
            layout.addLayout(button_layout)
            self.child_windows.append(users_window)
            users_window.show()
        except Exception as e:
            traceback_str = traceback.format_exc()
            self.logger.error(f"Error listing users: {e}")
            self.logger.error(f"Traceback: {traceback_str}")
            QMessageBox.critical(self.window, "Error", f"Failed to list users: {e}")

    def close_users_window(self):
        if hasattr(self, 'users_window'):
            self.users_window.close()
            self.users_window = None

    def sanitize(self):
        if not self.check_excel_loaded():
            return
        try:
            sanitized_file_path, _ = QFileDialog.getSaveFileName(self.window, "Save Sanitized Excel File", os.path.splitext(self.current_file_path)[0] + "_sanitized.xlsx", "Excel files (*.xlsx)")
            if not sanitized_file_path:
                return
            lock_acquired = False
            try:
                lock_status = self.acquire_file_lock(sanitized_file_path)
                if lock_status is None:  
                    return
                lock_acquired = True
                progress = QProgressBar()
                progress.setWindowTitle("Sanitizing File")
                progress.setGeometry(QRect(300, 300, 400, 30))
                progress.setWindowFlags(Qt.Window | Qt.WindowStaysOnTopHint)
                progress.setRange(0, len(self.current_workbook.sheetnames))
                progress.setValue(0)
                progress.show()
                ip_regex = re.compile(r'(\d{1,3}\.\d{1,3}\.\d{1,3})\.(\d{1,3})')
                http_regex = re.compile(r'(https?)(://)', re.IGNORECASE)
                domain_regex = re.compile(r'([a-zA-Z0-9][-a-zA-Z0-9]*\.[a-zA-Z0-9][-a-zA-Z0-9]*)\.([a-zA-Z]{2,})')
                for sheet_idx, sheet_name in enumerate(self.current_workbook.sheetnames):
                    progress.setValue(sheet_idx)
                    progress.setWindowTitle(f"Sanitizing: {sheet_name}")
                    QApplication.processEvents()
                    sheet = self.current_workbook[sheet_name]
                    for row in range(1, sheet.max_row + 1):
                        for col in range(1, sheet.max_column + 1):
                            cell = sheet.cell(row=row, column=col)
                            if cell.value and isinstance(cell.value, str):
                                cell.value = ip_regex.sub(r'\1[.]\2', cell.value)
                                cell.value = http_regex.sub(r'hxxp\2', cell.value)
                                cell.value = domain_regex.sub(r'\1[.]\2', cell.value)
                self.current_workbook.save(sanitized_file_path)
                progress.close()
                QMessageBox.information(self.window, "Sanitization Complete", f"File has been sanitized and saved to:\n{sanitized_file_path}"
                )
                reply = QMessageBox.question(self.window, "Load Sanitized File", "Do you want to load the sanitized file?", QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes)
                if reply == QMessageBox.Yes:
                    self.release_file_lock()
                    lock_acquired = False
                    self.load_data_into_treeview(file_path=sanitized_file_path)
            finally:
                if lock_acquired:
                    self.release_file_lock()
        except Exception as e:
            self.logger.error(f"Error during sanitization: {e}")
            traceback.print_exc()
            QMessageBox.critical(self.window, "Sanitization Error", f"An error occurred while sanitizing the file:\n{str(e)}")

    def __del__(self):
        self.logger.info("MainApp destructor called, cleaning up resources...")
        self.close_all_windows()
        
    def close_all_windows(self):
        if hasattr(self, 'child_windows'):
            for window in self.child_windows[:]:  
                try:
                    if (window and 
                        hasattr(window, 'isVisible') and 
                        isValid(window) and  
                        window.isVisible()):
                        window.close()
                except RuntimeError as e:
                    if "already deleted" in str(e):
                        self.logger.info(f"Window already deleted: {e}")
                    else:
                        self.logger.error(f"Error closing window: {e}")
                except Exception as e:
                    self.logger.error(f"Error closing window: {e}")
            self.child_windows.clear()

    def exit_application(self):
        self.close_all_windows()
        self.window.close()
        self.app.quit()

    def track_child_window(self, window):
        if window:
            if hasattr(window, 'setParent') and window.parent() is None:
                window.setParent(self.window, Qt.Window)
            self.child_windows.append(window)
            window.setAttribute(Qt.WA_DeleteOnClose, False)
            original_close = window.closeEvent if hasattr(window, 'closeEvent') else None
            
            def new_close_event(event):
                if window in self.child_windows:
                    self.child_windows.remove(window)
                if original_close:
                    original_close(event)
                else:
                    event.accept()
            window.closeEvent = new_close_event
            return window
        return None

    def application_cleanup(self):
        self.release_file_lock()
        for window in self.child_windows[:]:
            try:
                if window and isValid(window):  
                    self.logger.info(f"Final cleanup of window: {window.windowTitle() if hasattr(window, 'windowTitle') else 'Unknown'}")
                    window.setAttribute(Qt.WA_DeleteOnClose, True)
                    window.close()
            except RuntimeError as e:
                if "already deleted" in str(e):
                    self.logger.info(f"Window already deleted: {e}")
                else:
                    self.logger.error(f"Error in final cleanup: {e}")
            except Exception as e:
                self.logger.error(f"Error in final cleanup: {e}")


if __name__ == "__main__":
    main_app = MainApp()

    sys.exit(main_app.run())
